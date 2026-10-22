# -*- coding: utf-8 -*-
"""核心逻辑：读文件 / 从 JD 提词 / 匹配简历 / 打分 / 区分"只缺词"与"真缺经历"。

全部纯标准库，v1 不依赖模型，也不依赖第三方包。
读 PDF/DOCX 时才需要可选依赖（见 read_text）。
"""
import os
import re
import difflib

import kwlib

# ---------------------------------------------------------------------------
# 1. 读文件
# ---------------------------------------------------------------------------
TEXT_EXT = (".txt", ".md")


def read_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in TEXT_EXT:
        with open(path, encoding="utf-8") as f:
            return f.read()
    if ext in (".docx", ".docm"):
        try:
            import docx  # noqa: F401
        except ImportError:
            raise SystemExit(
                "读 .docx 需要 python-docx。在你的终端里跑：\n"
                "    pip install python-docx\n"
                "或把内容另存为 .txt 再传进来。"
            )
        import docx
        parts = [p.text for p in docx.Document(path).paragraphs]
        for t in docx.Document(path).tables:
            for row in t.rows:
                parts.append("\t".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext == ".pdf":
        try:
            import pdfplumber  # noqa: F401
        except ImportError:
            raise SystemExit(
                "读 .pdf 需要 pdfplumber。在你的终端里跑：\n"
                "    pip install pdfplumber\n"
                "或把内容另存为 .txt 再传进来。"
            )
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages)
    raise SystemExit("不支持的文件类型：%s（支持 .txt .md .docx .pdf）" % ext)


# ---------------------------------------------------------------------------
# 2. 从 JD 里找出"招聘方会搜的词"
# ---------------------------------------------------------------------------
def _norm(s):
    """归一化后再匹配。

    BOSS 直聘等平台的 JD 常见排版是「AI 工具」「需求 评审」——中文之间插了空格。
    精确子串匹配会因此全部漏掉，所以先把中文字符之间的空白去掉。
    英文词之间的空格必须保留（"Product Manager" 有意义）。
    """
    s = (s or "").lower()
    s = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", s)
    # 「B 端」「AI 工具」「Python 程序设计」这类 中英之间 的空格也要去掉，
    # 否则 "B端" 永远匹配不上 "B 端"。
    s = re.sub(r"(?<=[A-Za-z0-9])\s+(?=[\u4e00-\u9fff])", "", s)
    s = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[A-Za-z0-9])", "", s)
    s = re.sub(r"[ \t]+", " ", s)
    return s


def _term_of(x):
    """关键词既可能是 (layer, term) 元组（规则库），也可能是带 aliases 的 dict（v2）。"""
    return x["term"] if isinstance(x, dict) else x[1]


def dedupe_substrings(hits):
    """去掉被更长命中词包含的短词（避免"售前"和"售前技术支持"重复计数）。"""
    out = []
    for it in hits:
        low = _term_of(it).lower()
        swallowed = any(
            low != _term_of(o).lower() and low in _term_of(o).lower()
            for o in hits
        )
        if not swallowed:
            out.append(it)
    return out


def find_jd_keywords(jd_text):
    """返回 [(layer, term)]：库里出现且命中 JD 原文的词。"""
    t = _norm(jd_text)
    hits, seen = [], set()
    for layer in kwlib.LAYERS:
        for term in kwlib.KW.get(layer, []):
            key = term.lower()
            if key in seen:
                continue
            if key in t:
                seen.add(key)
                hits.append((layer, term))
    return dedupe_substrings(hits)


# ---------------------------------------------------------------------------
# 3. 匹配简历
# ---------------------------------------------------------------------------
def match_resume(resume_text, keywords):
    """对每个关键词，同时算出「精确命中」和「宽松命中」。

    精确命中 = 简历里出现了这个词本身（招聘方搜索时会字面匹配到）。
    宽松命中 = 没精确命中，但简历里有同一段经历的另一种说法（RELATED 里的锚点）。

    为什么要分开：中文构词很活。简历写「原型」「推进」「沟通」，JD 写「原型搭建」
    「项目推进」「沟通表达」，字面一个都对不上——但这些人的经历其实是在的。
    两个数的差额，就是「只换个说法就能拿到的分」。
    """
    r = _norm(resume_text)
    rows = []
    for it in keywords:
        if isinstance(it, dict):
            layer, term = it["layer"], it["term"]
            aliases = [a for a in (it.get("aliases") or []) if a]
            src = it.get("src", "rules")
        else:
            layer, term = it
            aliases = []
            src = "rules"
        cnt = r.count(_norm(term))
        strict = cnt > 0
        # 宽松命中的两个来源：模型给的 aliases（v2）+ 规则库的 RELATED 锚点（v1）
        found = [a for a in aliases if _norm(a) in r]
        for a in kwlib.RELATED.get(term, []):
            if _norm(a) in r and a not in found:
                found.append(a)
        rows.append({
            "layer": layer,
            "layer_cn": kwlib.LAYER_CN[layer],
            "term": term,
            "necessity": kwlib.NECESSITY[layer],
            "src": src,
            "hit": strict,
            "count": cnt,
            "hit_lenient": strict or bool(found),
            "anchors": found,
            "missing_kind": "hit" if strict else ("word_only" if found else "no_exp"),
        })
    return rows


# ---------------------------------------------------------------------------
# 4. 区分"只缺词"与"真缺经历"
# ---------------------------------------------------------------------------
def classify_missing(term, resume_text):
    """看简历里有没有"同一件经历、换了说法"的痕迹。

    返回 (kind, anchors)：
      kind = "word_only"  → 有相关经历，只是没用招聘方的词（可放心加）
      kind = "no_exp"     → 连相关痕迹都没有，硬加会在面试时崩（不要加）
    """
    anchors = kwlib.RELATED.get(term, [])
    r = _norm(resume_text)
    found = [a for a in anchors if a.lower() in r]
    return ("word_only" if found else "no_exp"), found


# ---------------------------------------------------------------------------
# 5. 打分
# ---------------------------------------------------------------------------
def _cov(rows, layers):
    sel = [x for x in rows if x["layer"] in layers]
    if not sel:
        return None
    return sum(1 for x in sel if x["hit"]) / len(sel) * 100.0


def score(rows):
    """返回 {"精确": {...}, "宽松": {...}} 两组分项覆盖率。

    精确 —— 现在这个简历，招聘方字面能不能搜到它说的东西。
    宽松 —— 把「同一段经历的另一种说法」也算上，即补齐措辞后的上限。
    """
    # 只用规则库那部分算分：模型词是 JD 原话长短语，字面几乎不会命中，
    # 放进分母会把覆盖率打成 0%，让数字失去区分度、也无法和历史对齐。
    base = [x for x in rows if x.get("src", "rules") == "rules"] or rows

    def cov(layers, key):
        sel = [x for x in base if x["layer"] in layers]
        if not sel:
            return None
        return sum(1 for x in sel if x[key]) / len(sel) * 100.0

    groups = (("title", "hard_skill"), ("title",), ("hard_skill",),
              ("tool", "domain"), ("soft", "bonus"))
    names = ("必需词", "职位名", "硬技能", "工具领域", "软性")

    out = {"精确": {}, "宽松": {}}
    for label, key in (("精确", "hit"), ("宽松", "hit_lenient")):
        for name, layers in zip(names, groups):
            out[label][name] = cov(layers, key)
    return out


# ---------------------------------------------------------------------------
# 6. 解析标题
# ---------------------------------------------------------------------------
def parse_resume_title(resume_text):
    for line in resume_text.splitlines():
        if "求职意向" in line:
            return line.split("：", 1)[-1].split(":", 1)[-1].strip()
    return ""


def parse_jd_title(jd_text, fallback=""):
    for line in jd_text.splitlines():
        m = re.match(r"^\s*岗位[:：]\s*(.+)$", line)
        if m:
            return m.group(1).strip()
    return fallback


def title_alignment(resume_title, jd_title):
    if not resume_title or not jd_title:
        return None
    return difflib.SequenceMatcher(
        None, resume_title.replace(" ", ""), jd_title.replace(" ", "")
    ).ratio() * 100.0


# ---------------------------------------------------------------------------
# 7. 一次跑完
# ---------------------------------------------------------------------------
def analyze(jd_text, resume_text, jd_name="", keywords=None):
    jd_title = parse_jd_title(jd_text, jd_name)
    resume_title = parse_resume_title(resume_text)
    if keywords is None:
        keywords = find_jd_keywords(jd_text)
    rows = match_resume(resume_text, keywords)
    return {
        "jd_name": jd_name,
        "jd_title": jd_title,
        "resume_title": resume_title,
        "title_align": title_alignment(resume_title, jd_title),
        "rows": rows,
        "scores": score(rows),
    }
